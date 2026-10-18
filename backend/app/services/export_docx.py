import io
import os
import re
from html.parser import HTMLParser
from typing import List, Optional

from docx import Document as DocxDocument
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls

from app.models.document import Document
from app.services.media_resolver import resolve_image_data
from app.utils.logger import logger

def set_cell_background(cell, color_hex: str):
    """Sets background shading of a docx table cell."""
    shading_elm = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{color_hex}"/>')
    cell._tc.get_or_add_tcPr().append(shading_elm)

def set_cell_left_border(cell, color_hex: str, size: int = 24):
    """Sets a left border on a cell for callouts."""
    tcPr = cell._tc.get_or_add_tcPr()
    tcBorders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>\n'
        f'  <w:left w:val="single" w:sz="{size}" w:space="0" w:color="{color_hex}"/>\n'
        f'  <w:top w:val="none"/>\n'
        f'  <w:right w:val="none"/>\n'
        f'  <w:bottom w:val="none"/>\n'
        f'</w:tcBorders>'
    )
    tcPr.append(tcBorders)

def render_image_in_docx(docx: DocxDocument, src: str, alt_text: str = ""):
    """Resolves and embeds an image centered into a docx document."""
    img_data = resolve_image_data(src)
    if not img_data:
        if alt_text:
            p = docx.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r = p.add_run(f"[Image: {alt_text}]")
            r.font.name = "Arial"
            r.font.size = Pt(8.5)
            r.font.italic = True
            r.font.color.rgb = RGBColor(148, 163, 184)
        return

    _, raw_bytes, _, w, h = img_data
    try:
        target_w = min(5.5, max(1.8, round(w / 110.0, 2)))
        p = docx.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(8)
        p.paragraph_format.space_after = Pt(4)
        run = p.add_run()
        run.add_picture(io.BytesIO(raw_bytes), width=Inches(target_w))
        if alt_text:
            cp = docx.add_paragraph()
            cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
            cp.paragraph_format.space_after = Pt(8)
            crun = cp.add_run(f"Figure: {alt_text}")
            crun.font.name = "Arial"
            crun.font.size = Pt(8.5)
            crun.font.italic = True
            crun.font.color.rgb = RGBColor(100, 116, 139)
    except Exception as e:
        logger.warning(f"Error embedding picture in docx: {e}")

class DocxHtmlParser(HTMLParser):
    def __init__(self, docx: DocxDocument):
        super().__init__()
        self.docx = docx
        self.current_p = None
        self.bold = False
        self.italic = False
        self.underline = False
        self.strike = False
        self.code_mode = False
        self.in_pre = False
        self.pre_buffer = []
        self.in_table = False
        self.table_rows = []
        self.current_row = []
        self.in_cell = False
        self.cell_buffer = []
        self.list_stack = []  # 'ul' or 'ol'
        self.in_blockquote = False
        self.quote_buffer = []

    def handle_starttag(self, tag, attrs):
        attrs_dict = dict(attrs)
        tag = tag.lower()

        if tag == "p":
            if not self.in_pre and not self.in_table and not self.in_blockquote:
                self.current_p = self.docx.add_paragraph()
                self.current_p.paragraph_format.space_after = Pt(6)

        elif tag in ("h1", "h2", "h3", "h4"):
            self.current_p = self.docx.add_paragraph()
            self.current_p.paragraph_format.space_before = Pt(12)
            self.current_p.paragraph_format.space_after = Pt(4)

        elif tag == "strong" or tag == "b":
            self.bold = True
        elif tag == "em" or tag == "i":
            self.italic = True
        elif tag == "u":
            self.underline = True
        elif tag in ("s", "strike", "del"):
            self.strike = True
        elif tag == "code" and not self.in_pre:
            self.code_mode = True

        elif tag == "ul":
            self.list_stack.append("ul")
        elif tag == "ol":
            self.list_stack.append("ol")
        elif tag == "li":
            list_style = 'List Bullet' if self.list_stack and self.list_stack[-1] == 'ul' else 'List Number'
            self.current_p = self.docx.add_paragraph(style=list_style)
            self.current_p.paragraph_format.space_after = Pt(3)

        elif tag == "blockquote":
            self.in_blockquote = True
            self.quote_buffer = []

        elif tag == "pre":
            self.in_pre = True
            self.pre_buffer = []

        elif tag == "table":
            self.in_table = True
            self.table_rows = []

        elif tag == "tr":
            self.current_row = []

        elif tag in ("th", "td"):
            self.in_cell = True
            self.cell_buffer = []

        elif tag == "img":
            src = attrs_dict.get("src", "")
            alt = attrs_dict.get("alt", "")
            if src:
                render_image_in_docx(self.docx, src, alt)

        elif tag == "br":
            if self.current_p:
                self.current_p.add_run("\n")

    def handle_endtag(self, tag):
        tag = tag.lower()
        if tag == "strong" or tag == "b":
            self.bold = False
        elif tag == "em" or tag == "i":
            self.italic = False
        elif tag == "u":
            self.underline = False
        elif tag in ("s", "strike", "del"):
            self.strike = False
        elif tag == "code" and not self.in_pre:
            self.code_mode = False

        elif tag in ("ul", "ol"):
            if self.list_stack:
                self.list_stack.pop()

        elif tag == "blockquote":
            self.in_blockquote = False
            q_text = "".join(self.quote_buffer).strip()
            if q_text:
                tbl = self.docx.add_table(rows=1, cols=1)
                tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
                cell = tbl.cell(0, 0)
                set_cell_background(cell, "EFF6FF")
                set_cell_left_border(cell, "3B82F6", size=20)
                cp = cell.paragraphs[0]
                cp.paragraph_format.space_before = Pt(4)
                cp.paragraph_format.space_after = Pt(4)
                r = cp.add_run(q_text)
                r.font.name = "Arial"
                r.font.size = Pt(9.5)
                r.font.color.rgb = RGBColor(30, 41, 59)
                self.docx.add_paragraph().paragraph_format.space_after = Pt(4)
            self.quote_buffer = []

        elif tag == "pre":
            self.in_pre = False
            code_text = "".join(self.pre_buffer).strip()
            if code_text:
                tbl = self.docx.add_table(rows=1, cols=1)
                tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
                cell = tbl.cell(0, 0)
                set_cell_background(cell, "F1F5F9")
                set_cell_left_border(cell, "6366F1", size=18)
                cp = cell.paragraphs[0]
                cp.paragraph_format.space_before = Pt(4)
                cp.paragraph_format.space_after = Pt(4)
                c_run = cp.add_run(code_text)
                c_run.font.name = "Courier New"
                c_run.font.size = Pt(9.5)
                c_run.font.color.rgb = RGBColor(15, 23, 42)
                self.docx.add_paragraph().paragraph_format.space_after = Pt(4)
            self.pre_buffer = []

        elif tag in ("th", "td"):
            self.in_cell = False
            self.current_row.append("".join(self.cell_buffer).strip())
            self.cell_buffer = []

        elif tag == "tr":
            if self.current_row:
                self.table_rows.append(self.current_row)
            self.current_row = []

        elif tag == "table":
            self.in_table = False
            if self.table_rows:
                num_cols = max(len(r) for r in self.table_rows)
                tbl = self.docx.add_table(rows=len(self.table_rows), cols=num_cols)
                tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
                for r_idx, row in enumerate(self.table_rows):
                    for c_idx in range(num_cols):
                        cell_text = row[c_idx] if c_idx < len(row) else ""
                        c = tbl.cell(r_idx, c_idx)
                        if r_idx == 0:
                            set_cell_background(c, "E2E8F0")
                        elif r_idx % 2 == 1:
                            set_cell_background(c, "F8FAFC")
                        p = c.paragraphs[0]
                        p.paragraph_format.space_before = Pt(3)
                        p.paragraph_format.space_after = Pt(3)
                        r = p.add_run(cell_text)
                        r.font.name = "Arial"
                        r.font.size = Pt(9.5)
                        if r_idx == 0:
                            r.font.bold = True
                self.docx.add_paragraph().paragraph_format.space_after = Pt(6)
            self.table_rows = []

        elif tag == "p":
            self.current_p = None

    def handle_data(self, data):
        if self.in_pre:
            self.pre_buffer.append(data)
            return

        if self.in_cell:
            self.cell_buffer.append(data)
            return

        if self.in_blockquote:
            self.quote_buffer.append(data)
            return

        text = data.strip() if not self.current_p else data
        if not text:
            return

        if self.current_p is None:
            self.current_p = self.docx.add_paragraph()
            self.current_p.paragraph_format.space_after = Pt(6)

        run = self.current_p.add_run(text)
        run.font.name = "Courier New" if self.code_mode else "Arial"
        run.font.size = Pt(9.5 if self.code_mode else 10)
        if self.bold:
            run.bold = True
        if self.italic:
            run.italic = True
        if self.underline:
            run.underline = True
        if self.strike:
            run.font.strike = True
        if self.code_mode:
            run.font.color.rgb = RGBColor(15, 23, 42)

def is_html_content(text: str) -> bool:
    if not text:
        return False
    return bool(re.search(r"<(p|div|table|h[1-6]|ul|ol|li|blockquote|pre|img|strong|em|span)[^>]*>", text, re.IGNORECASE))

def render_markdown_table_docx(docx: DocxDocument, table_lines: list):
    if not table_lines or len(table_lines) < 2:
        return
    parsed_rows = []
    for line in table_lines:
        line_clean = line.strip().strip("|")
        if re.match(r"^[\s\-:|]+$", line_clean):
            continue
        cells = [c.strip() for c in line_clean.split("|")]
        parsed_rows.append(cells)

    if not parsed_rows:
        return

    num_cols = max(len(r) for r in parsed_rows)
    tbl = docx.add_table(rows=len(parsed_rows), cols=num_cols)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER

    for r_idx, row in enumerate(parsed_rows):
        for c_idx in range(num_cols):
            cell_text = row[c_idx] if c_idx < len(row) else ""
            c = tbl.cell(r_idx, c_idx)
            if r_idx == 0:
                set_cell_background(c, "E2E8F0")
            elif r_idx % 2 == 1:
                set_cell_background(c, "F8FAFC")
            p = c.paragraphs[0]
            p.paragraph_format.space_before = Pt(3)
            p.paragraph_format.space_after = Pt(3)
            r = p.add_run(cell_text)
            r.font.name = "Arial"
            r.font.size = Pt(9.5)
            if r_idx == 0:
                r.font.bold = True
    docx.add_paragraph().paragraph_format.space_after = Pt(6)

def render_markdown_section_docx(docx: DocxDocument, content: str):
    lines = content.splitlines()
    table_buffer = []

    for line in lines:
        line_str = line.strip()

        # Check for table line
        if line_str.startswith("|") and line_str.endswith("|"):
            table_buffer.append(line_str)
            continue
        else:
            if table_buffer:
                render_markdown_table_docx(docx, table_buffer)
                table_buffer = []

        if not line_str:
            continue

        # HTML image tag embedded in markdown: <img ...>
        if "<img" in line_str.lower():
            src_m = re.search(r'src=["\']([^"\']+)["\']', line_str, re.IGNORECASE)
            alt_m = re.search(r'alt=["\']([^"\']*)["\']', line_str, re.IGNORECASE)
            if src_m:
                render_image_in_docx(docx, src_m.group(1), alt_m.group(1) if alt_m else "")
            continue

        # Image in markdown: ![alt](url)
        img_match = re.match(r"^!\[(.*?)\]\((.*?)\)$", line_str)
        if img_match:
            render_image_in_docx(docx, img_match.group(2), img_match.group(1))
            continue

        if line_str.startswith("#"):
            subh = docx.add_paragraph()
            subh.paragraph_format.space_before = Pt(10)
            subh.paragraph_format.space_after = Pt(4)
            subh_run = subh.add_run(re.sub(r"^#+\s*", "", line_str))
            subh_run.font.name = "Arial"
            subh_run.font.size = Pt(12)
            subh_run.font.bold = True
            subh_run.font.color.rgb = RGBColor(51, 65, 85)
        elif line_str.startswith("- ") or line_str.startswith("* "):
            bp = docx.add_paragraph(style='List Bullet')
            bp.paragraph_format.space_after = Pt(3)
            brun = bp.add_run(line_str[2:])
            brun.font.name = "Arial"
            brun.font.size = Pt(10)
        elif re.match(r"^\d+\.\s", line_str):
            np = docx.add_paragraph(style='List Number')
            np.paragraph_format.space_after = Pt(3)
            nrun = np.add_run(re.sub(r"^\d+\.\s", "", line_str))
            nrun.font.name = "Arial"
            nrun.font.size = Pt(10)
        elif line_str.startswith("> "):
            qp = docx.add_paragraph()
            qp.paragraph_format.left_indent = Inches(0.3)
            qp.paragraph_format.space_after = Pt(4)
            qrun = qp.add_run(line_str[2:])
            qrun.font.name = "Arial"
            qrun.font.italic = True
            qrun.font.size = Pt(9.5)
            qrun.font.color.rgb = RGBColor(71, 85, 105)
        else:
            p = docx.add_paragraph()
            p.paragraph_format.space_after = Pt(6)
            prun = p.add_run(line_str)
            prun.font.name = "Arial"
            prun.font.size = Pt(10)
            prun.font.color.rgb = RGBColor(30, 41, 59)

    if table_buffer:
        render_markdown_table_docx(docx, table_buffer)

def generate_docx_stream(doc: Document) -> io.BytesIO:
    docx = DocxDocument()

    # Set page margins (1 inch all sides)
    for section in docx.sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)

        # Add running header
        header = section.header
        hp = header.paragraphs[0]
        hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        hrun = hp.add_run(f"NoteForge AI  |  {doc.title}")
        hrun.font.name = "Arial"
        hrun.font.size = Pt(8.5)
        hrun.font.color.rgb = RGBColor(128, 128, 128)

        # Add running footer
        footer = section.footer
        fp = footer.paragraphs[0]
        fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        frun = fp.add_run(f"Version {doc.current_version}  •  Study Notes Documentation")
        frun.font.name = "Arial"
        frun.font.size = Pt(8.5)
        frun.font.color.rgb = RGBColor(128, 128, 128)

    # Document Title
    title_p = docx.add_paragraph()
    title_p.paragraph_format.space_before = Pt(12)
    title_p.paragraph_format.space_after = Pt(4)
    run_title = title_p.add_run(doc.title)
    run_title.font.name = "Arial"
    run_title.font.size = Pt(24)
    run_title.font.bold = True
    run_title.font.color.rgb = RGBColor(30, 41, 59)

    # Subject & Metadata Subtitle
    sub_p = docx.add_paragraph()
    sub_p.paragraph_format.space_after = Pt(16)
    sub_run = sub_p.add_run(f"Subject: {doc.subject or 'General'}   •   Version: {doc.current_version}")
    sub_run.font.name = "Arial"
    sub_run.font.size = Pt(11)
    sub_run.font.color.rgb = RGBColor(100, 116, 139)

    # Summary
    if doc.summary:
        sum_p = docx.add_paragraph()
        sum_p.paragraph_format.space_after = Pt(18)
        sum_run = sum_p.add_run(doc.summary)
        sum_run.font.name = "Arial"
        sum_run.font.size = Pt(10.5)
        sum_run.font.italic = True
        sum_run.font.color.rgb = RGBColor(71, 85, 105)

    # Divider line
    divider = docx.add_paragraph()
    divider.paragraph_format.space_after = Pt(16)
    d_run = divider.add_run("____________________________________________________________________")
    d_run.font.color.rgb = RGBColor(226, 232, 240)

    # Document Sections
    for s_idx, section in enumerate(doc.sections):
        # Section Heading
        h2 = docx.add_paragraph()
        h2.paragraph_format.space_before = Pt(18)
        h2.paragraph_format.space_after = Pt(6)
        h2_run = h2.add_run(f"{s_idx + 1}. {section.title}")
        h2_run.font.name = "Arial"
        h2_run.font.size = Pt(15)
        h2_run.font.bold = True
        h2_run.font.color.rgb = RGBColor(15, 23, 42)

        # Body: check if HTML or Markdown
        if section.content_markdown:
            content = section.content_markdown.strip()
            if is_html_content(content):
                parser = DocxHtmlParser(docx)
                parser.feed(content)
            else:
                render_markdown_section_docx(docx, content)

        # Structured Blocks (Code blocks, callouts, tables, images)
        for block in (section.structured_blocks or []):
            b_type = block.get("type") if isinstance(block, dict) else block.type
            b_data = block.get("data", {}) if isinstance(block, dict) else getattr(block, "data", {})

            if b_type == "code":
                lang = b_data.get("language", "code").upper()
                code_text = b_data.get("code", "")

                tbl = docx.add_table(rows=1, cols=1)
                tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
                cell = tbl.cell(0, 0)
                set_cell_background(cell, "F1F5F9")
                set_cell_left_border(cell, "6366F1", size=18)

                cp = cell.paragraphs[0]
                cp.paragraph_format.space_before = Pt(4)
                cp.paragraph_format.space_after = Pt(4)

                lang_run = cp.add_run(f"[{lang}]\n")
                lang_run.font.name = "Courier New"
                lang_run.font.size = Pt(8.5)
                lang_run.font.bold = True
                lang_run.font.color.rgb = RGBColor(99, 102, 241)

                c_run = cp.add_run(code_text)
                c_run.font.name = "Courier New"
                c_run.font.size = Pt(9.5)
                c_run.font.color.rgb = RGBColor(15, 23, 42)

                docx.add_paragraph().paragraph_format.space_after = Pt(6)

            elif b_type == "callout":
                c_type = b_data.get("callout_type", "important")
                c_title = b_data.get("title", "Note")
                c_body = b_data.get("body", "")

                color_map = {
                    "warning": ("FEF2F2", "EF4444"),
                    "important": ("EFF6FF", "3B82F6"),
                    "tip": ("F0FDF4", "10B981"),
                    "info": ("F8FAFC", "64748B")
                }
                bg_color, border_color = color_map.get(c_type, ("EFF6FF", "3B82F6"))

                tbl = docx.add_table(rows=1, cols=1)
                tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
                cell = tbl.cell(0, 0)
                set_cell_background(cell, bg_color)
                set_cell_left_border(cell, border_color, size=24)

                cp = cell.paragraphs[0]
                cp.paragraph_format.space_before = Pt(6)
                cp.paragraph_format.space_after = Pt(6)

                trun = cp.add_run(f"{c_title.upper()}: ")
                trun.font.name = "Arial"
                trun.font.size = Pt(9.5)
                trun.font.bold = True

                brun = cp.add_run(c_body)
                brun.font.name = "Arial"
                brun.font.size = Pt(9.5)

                docx.add_paragraph().paragraph_format.space_after = Pt(6)

            elif b_type == "table":
                headers = b_data.get("headers", [])
                rows = b_data.get("rows", [])
                if headers and rows:
                    tbl = docx.add_table(rows=len(rows) + 1, cols=len(headers))
                    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
                    for c_i, header_txt in enumerate(headers):
                        c = tbl.cell(0, c_i)
                        set_cell_background(c, "E2E8F0")
                        p = c.paragraphs[0]
                        p.paragraph_format.space_before = Pt(3)
                        p.paragraph_format.space_after = Pt(3)
                        r = p.add_run(str(header_txt))
                        r.font.name = "Arial"
                        r.font.size = Pt(9.5)
                        r.font.bold = True

                    for r_i, row in enumerate(rows):
                        for c_i, val in enumerate(row):
                            c = tbl.cell(r_i + 1, c_i)
                            if r_i % 2 == 1:
                                set_cell_background(c, "F8FAFC")
                            p = c.paragraphs[0]
                            p.paragraph_format.space_before = Pt(3)
                            p.paragraph_format.space_after = Pt(3)
                            r = p.add_run(str(val))
                            r.font.name = "Arial"
                            r.font.size = Pt(9.5)

                    docx.add_paragraph().paragraph_format.space_after = Pt(6)

            elif b_type == "image":
                src_url = b_data.get("url", "")
                alt_text = b_data.get("caption", "")
                render_image_in_docx(docx, src_url, alt_text)

    # Save to memory stream
    output_stream = io.BytesIO()
    docx.save(output_stream)
    output_stream.seek(0)
    return output_stream
